from __future__ import annotations

import json
from pathlib import Path

from PIL import Image, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parents[1]
STUDY = json.loads((ROOT / "content/sound-xr-image-study.json").read_text(encoding="utf-8"))
IMAGE_SOURCES = [
    Path("/Users/taebin/Pictures/Photos Library.photoslibrary/resources/derivatives/2/2C3BD736-0284-4476-A8A6-811DCB4A554A_1_105_c.jpeg"),
    Path("/Users/taebin/Pictures/Photos Library.photoslibrary/resources/derivatives/2/25C352C0-7040-4F65-BCCF-F7DBE7D18B6F_1_102_o.jpeg"),
]
IMAGE_DIR = ROOT / "public/framer/projects/spatial-renderer-system-identification"
REPORT_DIR = ROOT / "public/reports"
KOREAN_FONT = Path("/System/Library/AssetsV2/com_apple_MobileAsset_Font7/bad9b4bf17cf1669dde54184ba4431c22dcad27b.asset/AssetData/NanumGothic.ttc")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def set_style_font(style, name: str) -> None:
    style.font.name = name
    r_fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        r_fonts.attrib.pop(qn(f"w:{attribute}"), None)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), name)


def set_run_font(run, name: str) -> None:
    run.font.name = name
    r_fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        r_fonts.attrib.pop(qn(f"w:{attribute}"), None)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{attribute}"), name)


def apply_explicit_fonts(doc: Document, locale: str) -> None:
    body_font = "Arial Unicode MS" if locale == "ko" else "Arial"
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in doc.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            set_run_font(run, body_font)


def configure_document(doc: Document, locale: str) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.35)
    section.bottom_margin = Cm(2.15)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)

    body_font = "Arial Unicode MS" if locale == "ko" else "Arial"
    normal = doc.styles["Normal"]
    set_style_font(normal, body_font)
    normal.font.size = Pt(9.7)
    normal.font.color.rgb = RGBColor(30, 30, 30)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.28

    title = doc.styles["Title"]
    set_style_font(title, body_font)
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(10)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12)):
        style = doc.styles[style_name]
        set_style_font(style, body_font)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = "SOUND xR IMAGE  /  RENDERER SYSTEM IDENTIFICATION"
    header.style = doc.styles["Normal"]
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.color.rgb = RGBColor(105, 105, 105)
    header.runs[0].font.bold = True
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(105, 105, 105)


def add_findings_table(doc: Document, locale: str) -> None:
    if locale == "ko":
        headers = ["항목", "주요 결과", "해석 범위"]
        rows = [
            ["정상 출력", "대표 잔차 약 0.005 dB", "거리 가중치와 전력 정규화를 지지"],
            ["Precision", "거리 두 배당 약 8–12 dB", "정규화 전 후보 계수"],
            ["Size", "ON 최대 0.0047 dB  OFF 최대 0.0051 dB", "시험 배치와 Object Z=0 범위"],
            ["Speaker ID", "배정별 최대 0.0128 dB와 1.1728 dB", "좌표만으로 설명되지 않음"],
            ["경계와 Z", "일부 투영 적합  Z 스윕 변화 없음", "선분 선택과 XY OFF 검증 필요"],
        ]
    else:
        headers = ["Topic", "Primary result", "Interpretation range"]
        rows = [
            ["Normal output", "Representative residual about 0.005 dB", "Supports distance weighting and power normalization"],
            ["Precision", "About 8–12 dB per distance doubling", "Candidate coefficient before normalization"],
            ["Size", "ON max 0.0047 dB  OFF max 0.0051 dB", "Tested layouts and Object Z=0 only"],
            ["Speaker ID", "Assignment maxima 0.0128 dB and 1.1728 dB", "Not explained by coordinates alone"],
            ["Boundary and Z", "Some projections fit  Z sweep unchanged", "Segment rule and XY OFF remain open"],
        ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(3.0), Cm(5.0), Cm(7.0)]
    header = table.rows[0]
    for idx, (cell, text) in enumerate(zip(header.cells, headers)):
        cell.width = widths[idx]
        cell.text = text
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "253347")
        set_cell_margins(cell, top=80, bottom=80)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for idx, (cell, text) in enumerate(zip(row.cells, values)):
            cell.width = widths[idx]
            cell.text = text
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=75, bottom=75)
            if row_index % 2:
                set_cell_shading(cell, "F2F5F8")
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(8)


def build_report(locale: str, output: Path) -> None:
    report = STUDY[locale]
    meta = STUDY["metadata"]
    doc = Document()
    configure_document(doc, locale)
    doc.core_properties.title = report["title"]
    doc.core_properties.subject = report["subtitle"]
    doc.core_properties.author = "Taebin Yoo"

    title = doc.add_paragraph(report["title"], style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = doc.add_paragraph(report["subtitle"])
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.color.rgb = RGBColor(85, 85, 85)
    subtitle.paragraph_format.space_after = Pt(16)
    meta_line = doc.add_paragraph(
        (f"결과보고서  |  {meta['date']}  |  개정 {meta['revision']}" if locale == "ko" else f"Research report  |  {meta['date']}  |  Revision {meta['revision']}")
    )
    meta_line.runs[0].font.size = Pt(8)
    meta_line.runs[0].font.bold = True
    meta_line.runs[0].font.color.rgb = RGBColor(107, 107, 107)
    meta_line.paragraph_format.space_after = Pt(20)

    doc.add_heading("초록" if locale == "ko" else "Abstract", level=1)
    doc.add_paragraph(report["abstract"])
    keywords = doc.add_paragraph()
    lead = keywords.add_run("주요어  " if locale == "ko" else "Keywords  ")
    lead.bold = True
    keywords.add_run(" · ".join(report["keywords"]))
    keywords.paragraph_format.space_after = Pt(16)

    for section in report["sections"]:
        doc.add_heading(section["heading"], level=1)
        for paragraph in section["paragraphs"]:
            doc.add_paragraph(paragraph)
        if section.get("equations"):
            for equation in section["equations"]:
                p = doc.add_paragraph(equation)
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(7)
                p.paragraph_format.space_after = Pt(7)
                for run in p.runs:
                    run.font.name = "Cambria Math"
                    run.font.size = Pt(13)
            note = doc.add_paragraph(
                "D는 계산 위치와 스피커 사이 거리이며 a는 Precision에 따른 후보 거리 지수다." if locale == "ko" else "D is the distance from the calculated position to each loudspeaker, and a is the candidate Precision dependent distance exponent."
            )
            note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in note.runs:
                run.font.size = Pt(8)
                run.font.italic = True
                run.font.color.rgb = RGBColor(90, 90, 90)
        if section["heading"].startswith("4 "):
            add_findings_table(doc, locale)

    doc.add_heading("데이터 공개" if locale == "ko" else "Data Availability", level=1)
    doc.add_paragraph(
        "원본 측정 기록은 웹사이트의 실험 데이터 미리보기와 XLSX 파일로 제공한다. 후보식의 해석은 본 보고서에 명시한 조건 범위로 제한한다."
        if locale == "ko"
        else "The source measurements are provided through the website data preview and the XLSX workbook. Interpretation of the candidate equations is limited to the conditions stated in this report."
    )
    apply_explicit_fonts(doc, locale)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def build_images() -> None:
    IMAGE_DIR.mkdir(parents=True, exist_ok=True)
    for index, source in enumerate(IMAGE_SOURCES, start=1):
        with Image.open(source) as original:
            image = ImageOps.exif_transpose(original).convert("RGB")
            image.thumbnail((1800, 1800), Image.Resampling.LANCZOS)
            image.save(IMAGE_DIR / f"{index:02d}.jpeg", "JPEG", quality=88, optimize=True, progressive=True)
    old_cover = IMAGE_DIR / "cover.svg"
    if old_cover.exists():
        old_cover.unlink()


def ascii_hyphens(text: str) -> str:
    return text.replace("–", "-").replace("—", "-").replace("−", "-")


def pdf_findings(locale: str):
    if locale == "ko":
        return [
            ["항목", "주요 결과", "해석 범위"],
            ["정상 출력", "대표 잔차 약 0.005 dB", "거리 가중치와 전력 정규화를 지지"],
            ["Precision", "거리 두 배당 약 8-12 dB", "정규화 전 후보 계수"],
            ["Size", "ON 최대 0.0047 dB\nOFF 최대 0.0051 dB", "시험 배치와 Object Z=0 범위"],
            ["Speaker ID", "배정별 최대 0.0128 dB와 1.1728 dB", "좌표만으로 설명되지 않음"],
            ["경계와 Z", "일부 투영 적합\nZ 스윕 변화 없음", "선분 선택과 XY OFF 검증 필요"],
        ]
    return [
        ["Topic", "Primary result", "Interpretation range"],
        ["Normal output", "Representative residual about 0.005 dB", "Supports distance weighting and power normalization"],
        ["Precision", "About 8-12 dB per distance doubling", "Candidate coefficient before normalization"],
        ["Size", "ON max 0.0047 dB\nOFF max 0.0051 dB", "Tested layouts and Object Z=0 only"],
        ["Speaker ID", "Assignment maxima 0.0128 dB and 1.1728 dB", "Not explained by coordinates alone"],
        ["Boundary and Z", "Some projections fit\nZ sweep unchanged", "Segment rule and XY OFF remain open"],
    ]


def build_pdf(locale: str, output: Path) -> None:
    report = STUDY[locale]
    meta = STUDY["metadata"]
    if locale == "ko":
        pdfmetrics.registerFont(TTFont("NanumGothic", str(KOREAN_FONT), subfontIndex=0))
        regular = bold = "NanumGothic"
    else:
        regular, bold = "Helvetica", "Helvetica-Bold"

    doc = SimpleDocTemplate(
        str(output), pagesize=A4, leftMargin=23*mm, rightMargin=23*mm,
        topMargin=24*mm, bottomMargin=20*mm, title=report["title"], author="Taebin Yoo"
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("ReportTitle", parent=styles["Title"], fontName=bold, fontSize=24, leading=29, textColor=colors.black, alignment=TA_LEFT, spaceAfter=10)
    subtitle_style = ParagraphStyle("ReportSubtitle", parent=styles["Normal"], fontName=regular, fontSize=11, leading=16, textColor=colors.HexColor("#555555"), spaceAfter=11)
    meta_style = ParagraphStyle("ReportMeta", parent=styles["Normal"], fontName=bold, fontSize=7.5, leading=10, textColor=colors.HexColor("#6b6b6b"), spaceAfter=17)
    heading_style = ParagraphStyle("ReportHeading", parent=styles["Heading1"], fontName=bold, fontSize=15, leading=19, textColor=colors.black, spaceBefore=13, spaceAfter=7, keepWithNext=True)
    body_style = ParagraphStyle("ReportBody", parent=styles["BodyText"], fontName=regular, fontSize=9.3, leading=13.7, textColor=colors.HexColor("#202020"), spaceAfter=7)
    equation_style = ParagraphStyle("ReportEquation", parent=body_style, fontName=regular, fontSize=13, leading=18, alignment=TA_CENTER, spaceBefore=5, spaceAfter=5)
    small_style = ParagraphStyle("ReportSmall", parent=body_style, fontSize=7.6, leading=10.5, textColor=colors.HexColor("#626262"))
    table_header_style = ParagraphStyle("TableHeader", parent=small_style, fontName=bold, textColor=colors.white, alignment=TA_CENTER)
    table_cell_style = ParagraphStyle("TableCell", parent=small_style, textColor=colors.HexColor("#202020"), alignment=TA_LEFT)

    story = [
        Paragraph(ascii_hyphens(report["title"]), title_style),
        Paragraph(ascii_hyphens(report["subtitle"]), subtitle_style),
        Paragraph((f"결과보고서  |  {meta['date']}  |  개정 {meta['revision']}" if locale == "ko" else f"Research report  |  {meta['date']}  |  Revision {meta['revision']}"), meta_style),
        Paragraph("초록" if locale == "ko" else "Abstract", heading_style),
        Paragraph(ascii_hyphens(report["abstract"]), body_style),
        Paragraph(("주요어  " if locale == "ko" else "Keywords  ") + " · ".join(report["keywords"]), small_style),
        Spacer(1, 4*mm),
    ]
    for section in report["sections"]:
        story.append(Paragraph(ascii_hyphens(section["heading"]), heading_style))
        for paragraph in section["paragraphs"]:
            story.append(Paragraph(ascii_hyphens(paragraph), body_style))
        for equation in section.get("equations", []):
            story.append(Paragraph(ascii_hyphens(equation), equation_style))
        if section.get("equations"):
            note = "D는 계산 위치와 스피커 사이 거리이며 a는 Precision에 따른 후보 거리 지수다." if locale == "ko" else "D is the distance from the calculated position to each loudspeaker, and a is the candidate Precision dependent distance exponent."
            story.append(Paragraph(note, small_style))
        if section["heading"].startswith("4 "):
            raw = pdf_findings(locale)
            table_data = []
            for row_index, row in enumerate(raw):
                style = table_header_style if row_index == 0 else table_cell_style
                table_data.append([Paragraph(ascii_hyphens(value).replace("\n", "<br/>"), style) for value in row])
            table = Table(table_data, colWidths=[31*mm, 58*mm, 66*mm], repeatRows=1, hAlign="CENTER")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#253347")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d9d9d9")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5f8")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.extend([Spacer(1, 4*mm), table, Spacer(1, 3*mm)])
    story.append(Paragraph("데이터 공개" if locale == "ko" else "Data Availability", heading_style))
    availability = "원본 측정 기록은 웹사이트의 실험 데이터 미리보기와 XLSX 파일로 제공한다. 후보식의 해석은 본 보고서에 명시한 조건 범위로 제한한다." if locale == "ko" else "The source measurements are provided through the website data preview and the XLSX workbook. Interpretation of the candidate equations is limited to the conditions stated in this report."
    story.append(Paragraph(availability, body_style))

    def decorate_page(canvas, document):
        canvas.saveState()
        canvas.setFont(regular, 7)
        canvas.setFillColor(colors.HexColor("#707070"))
        canvas.drawString(23*mm, A4[1]-14*mm, "SOUND xR IMAGE  /  RENDERER SYSTEM IDENTIFICATION")
        canvas.drawRightString(A4[0]-23*mm, 11*mm, str(document.page))
        canvas.restoreState()

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.build(story, onFirstPage=decorate_page, onLaterPages=decorate_page)


def main() -> None:
    build_images()
    build_pdf("ko", REPORT_DIR / "sound-xr-image-renderer-report-ko.pdf")
    build_pdf("en", REPORT_DIR / "sound-xr-image-renderer-report-en.pdf")
    for stale_docx in (REPORT_DIR / "sound-xr-image-renderer-report-ko.docx", REPORT_DIR / "sound-xr-image-renderer-report-en.docx"):
        if stale_docx.exists():
            stale_docx.unlink()
    old_report = REPORT_DIR / "afc-renderer-analysis-ko.docx"
    if old_report.exists():
        old_report.unlink()


if __name__ == "__main__":
    main()
