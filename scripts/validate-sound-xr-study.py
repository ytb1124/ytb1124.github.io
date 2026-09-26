"""Recompute selected manuscript results from the original workbook (openpyxl).

Run with an optional workbook path; defaults to the published source workbook.
This checks recorded gains, not instrument accuracy or independent validation.
"""
import ast
import json
import math
import re
import sys
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
path=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'public/downloads/afc-renderer-experiment.xlsx'
s=openpyxl.load_workbook(path,data_only=True).active
tri=[(-5,0),(5,0),(0,8.66)]

def gains(d2,b=10):
    a=b/(20*math.log10(2))
    weights=[d**(-a/2) for d in d2]
    norm=math.sqrt(sum(w*w for w in weights))
    return [20*math.log10(w/norm) for w in weights]

def dbap(speakers,obj,b=10,width=0):
    return gains([sum((x-y)**2 for x,y in zip(sp,obj))+width**2 for sp in speakers],b)

def measured(row,start,count):return [s.cell(row,start+c).value for c in range(count)]
def project(o,a,b):
    v=[y-x for x,y in zip(a,b)]
    t=max(0,min(1,sum((x-y)*z for x,y,z in zip(o,a,v))/sum(z*z for z in v)))
    return [x+t*y for x,y in zip(a,v)]
results={}
def check(name,pairs,expected_max=None):
    errors=[m-p for ms,ps in pairs for m,p in zip(ms,ps) if isinstance(m,(float,int))]
    result={'records':len(pairs),'n':len(errors),'rms':math.sqrt(sum(e*e for e in errors)/len(errors)),'max':max(map(abs,errors))}
    if expected_max is not None:assert abs(result['max']-expected_max)<.000051,(name,result)
    results[name]=result

check('internal_x',[(measured(r,2,3),dbap(tri,(s.cell(r,1).value,2.89))) for r in range(2,9)],.0046)
check('internal_xy',[(measured(r,2,3),dbap(tri,ast.literal_eval(s.cell(r,1).value))) for r in range(20,25)],.0045)
check('precision',[(measured(r,2,3),dbap(tri,(2,2.89),b=r-87)) for r in range(95,100)],.0048)
check('width',[(measured(r,2,3),dbap(tri,(3,2),width=s.cell(r,1).value)) for r in range(46,55)],.0052)
pairs=[]
for r in range(294,307):
    w,h=map(float,re.findall(r'[\d.]+',s.cell(r,1).value))
    we=max(w,min(h,1))
    pairs.append((measured(r,2,4),gains([29+we*we,29+we*we,9+we*we,9+we*we])))
check('size_xy_on',pairs,.0046)
pairs=[]
for r in range(294,298):
    w,h=map(float,re.findall(r'[\d.]+',s.cell(r,10).value))
    if w+h:
        u=min(max(w,h),1);we=max(w,u);he=max(h,u)
        d2=[29+we*we,29+we*we,9+we*we+(10*we/he)**2,9+we*we+(10*we/he)**2]
    else:d2=[29,29,109,109]
    pairs.append((measured(r,11,4),gains(d2)))
check('size_xy_off',pairs,.0051)
sp=[(-4,-3),(5,-1),(-2,4),(3,6),(0,1)]
for name,pts,col in [('N',sp,2),('T',[sp[0],sp[1],sp[3],sp[2],sp[4]],12)]:
    check(name+'_direct',[(measured(359,col,5),dbap(pts,(1,3)))],1.1728 if name=='N' else .0128)
    check(name+'_prior_segment',[(measured(359,col,5),dbap(pts,project((1,3),pts[1],pts[2])))],9.2876 if name=='N' else 8.6677)
check('N_posthoc_segment',[(measured(359,2,5),dbap(sp,project((1,3),sp[0],sp[3])))],.0064)
assert all(measured(r,2,4)==[-5.48,-5.48,-6.63,-6.63] for r in range(83,92))
for start in (2,13):
    assert measured(320,start,6)==measured(322,start,6)
    assert measured(321,start,6)==measured(323,start,6)
results['da_paired_changes']={}
for r in range(310,316):
    base=measured(r,18,3)
    changes=[v-b for v,b in zip(measured(r,2,3),base) if isinstance(v,(float,int)) and isinstance(b,(float,int))]
    assert max(changes)-min(changes)<.010001
    results['da_paired_changes'][str(s.cell(r,1).value)]=[round(x,2) for x in changes]
print(json.dumps(results,indent=2))
